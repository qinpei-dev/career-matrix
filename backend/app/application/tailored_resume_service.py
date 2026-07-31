"""Evidence-gated tailored resume workflow."""

from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
from datetime import datetime, timedelta, timezone
from typing import Any

from docx import Document as DocxDocument
from sqlalchemy.orm import Session

from .crud_service import ResourceConflictError, ResourceNotFoundError
from .retrieval_service import RetrievalService
from ..core.security import redact_sensitive_text
from ..infrastructure.database.models import TailoredResume, TailoredResumeStatus
from ..infrastructure.database.repositories import (
    DocumentRepository,
    JobRepository,
    TailoredResumeRepository,
    UserRepository,
)
from ..infrastructure.llm.provider import ResumeTailoringProvider
from ..schemas.tailored_resume import TailoredResumeUpdate

MAX_EVIDENCE = 40
MAX_KEYWORDS = 30
GENERATION_STALE_AFTER = timedelta(minutes=10)
TOKEN_RE = re.compile(r"[A-Za-z][A-Za-z0-9+#.-]{1,30}|[\u4e00-\u9fff]{2,8}")
STOP_WORDS = {
    "and", "the", "with", "for", "from", "this", "that", "you", "your",
    "all", "need", "ignore", "rules", "invent", "years", "experience",
    "output", "key", "automatically", "apply", "send", "message", "of", "at",
    "工作", "岗位", "负责", "要求", "能力", "经验", "相关", "优先", "以及", "我们",
}


class UnsupportedTailoredContentError(ValueError):
    """Raised when an edit is not provably present in the source resume."""


class TailoredResumeService:
    def __init__(
        self,
        session: Session,
        user_email: str,
        provider: ResumeTailoringProvider,
        retrieval_service: RetrievalService,
    ) -> None:
        self.session = session
        self.user_email = user_email.strip().lower()
        self.provider = provider
        self.retrieval_service = retrieval_service
        self.users = UserRepository(session)
        self.jobs = JobRepository(session)
        self.documents = DocumentRepository(session)
        self.resumes = TailoredResumeRepository(session)

    def _user(self):
        return self.users.get_or_create_by_email(self.user_email)

    def _get(self, resume_id: uuid.UUID) -> TailoredResume:
        user = self._user()
        resume = self.resumes.get_for_user(resume_id, user.id)
        if resume is None:
            self.session.rollback()
            raise ResourceNotFoundError("tailored resume not found")
        return resume

    def create(self, job_id: uuid.UUID, source_document_id: uuid.UUID) -> TailoredResume:
        user = self._user()
        job = self.jobs.get_for_user(job_id, user.id)
        if job is None:
            self.session.rollback()
            raise ResourceNotFoundError("job not found")
        source = self.documents.get_for_user(source_document_id, user.id)
        if source is None:
            self.session.rollback()
            raise ResourceNotFoundError("document not found")
        if source.status != "ready":
            self.session.rollback()
            raise ResourceConflictError("source document is not ready")
        resume = self.resumes.add(
            TailoredResume(
                user_id=user.id,
                job_id=job.id,
                source_document_id=source.id,
                title=f"{job.title} - 定制简历",
            )
        )
        self.session.commit()
        self.session.refresh(resume)
        return resume

    def get(self, resume_id: uuid.UUID) -> TailoredResume:
        resume = self._get(resume_id)
        self.session.commit()
        return resume

    def list(self, job_id: uuid.UUID | None = None) -> list[TailoredResume]:
        user = self._user()
        if job_id is not None and self.jobs.get_for_user(job_id, user.id) is None:
            self.session.rollback()
            raise ResourceNotFoundError("job not found")
        result = self.resumes.list_for_user(user.id, job_id)
        self.session.commit()
        return result

    @staticmethod
    def _category(section: str) -> str:
        key = section.casefold()
        if any(word in key for word in ("project", "项目")):
            return "projects"
        if any(word in key for word in ("education", "学历", "教育")):
            return "education"
        if any(word in key for word in ("skill", "技能", "技术")):
            return "skills"
        if any(word in key for word in ("summary", "profile", "简介", "概述", "优势")):
            return "summary"
        return "experience"

    @staticmethod
    def _keywords(text: str) -> list[str]:
        unique: list[str] = []
        seen: set[str] = set()
        for token in TOKEN_RE.findall(text):
            key = token.casefold()
            if key in STOP_WORDS or key in seen:
                continue
            seen.add(key)
            unique.append(token)
        return unique[:MAX_KEYWORDS]

    @staticmethod
    def _parse_proposal(raw: str) -> dict[str, Any]:
        try:
            value = json.loads(raw)
        except (json.JSONDecodeError, TypeError) as exc:
            raise ValueError("model returned invalid tailoring JSON") from exc
        if not isinstance(value, dict):
            raise ValueError("model returned invalid tailoring JSON")
        return value

    def generate(self, resume_id: uuid.UUID) -> TailoredResume:
        resume = self._get(resume_id)
        if resume.status in (
            TailoredResumeStatus.GENERATED,
            TailoredResumeStatus.EDITING,
            TailoredResumeStatus.FINALIZED,
        ):
            self.session.commit()
            return resume
        if resume.is_generating:
            now = datetime.now(timezone.utc)
            self.resumes.recover_stale_generation(
                resume.id,
                resume.user_id,
                stale_before=now - GENERATION_STALE_AFTER,
                recovered_at=now,
            )
            resume = self._get(resume_id)
            if resume.is_generating:
                self.session.rollback()
                raise ResourceConflictError("tailored resume generation is already running")
        claimed = self.resumes.claim_generation(resume.id, resume.user_id, resume.version)
        if not claimed:
            raise ResourceConflictError("tailored resume generation is already running")

        try:
            resume = self._get(resume_id)
            job = self.jobs.get_for_user(resume.job_id, resume.user_id)
            source = self.documents.get_for_user(
                resume.source_document_id, resume.user_id
            )
            if job is None or source is None:
                raise ResourceNotFoundError("source resource not found")
            chunks = self.documents.list_chunks_for_document(source.id)
            if source.status != "ready" or not chunks:
                raise ResourceConflictError("source document has no ready evidence")

            retrieved = self.retrieval_service.search(
                f"{job.title}\n{job.description}", min(MAX_EVIDENCE, len(chunks)),
                document_id=source.id,
            )
            score_by_id = {str(item.chunk_id): item.score for item in retrieved}
            ordered_chunks = sorted(
                chunks,
                key=lambda item: (
                    score_by_id.get(str(item.id), -1.0),
                    -item.chunk_index,
                ),
                reverse=True,
            )[:MAX_EVIDENCE]
            evidence = [
                {
                    "id": str(chunk.id),
                    "source_type": "resume_chunk",
                    "source_id": str(chunk.id),
                    "document_id": str(source.id),
                    "section": redact_sensitive_text(chunk.section),
                    "content": redact_sensitive_text(chunk.content),
                    "rag_score": score_by_id.get(str(chunk.id)),
                }
                for chunk in ordered_chunks
            ]
            evidence_payload = json.dumps(evidence, ensure_ascii=False)
            proposal = self._parse_proposal(
                self.provider.tailor_resume(
                    job.title, job.description, evidence_payload
                )
            )

            by_id = {item["id"]: item for item in evidence}
            proposed_order = proposal.get("evidence_order", [])
            valid_order = [
                item for item in proposed_order
                if isinstance(item, str) and item in by_id
            ]
            remaining = [item["id"] for item in evidence if item["id"] not in valid_order]
            ordered_ids = valid_order + remaining
            ordered_evidence = [by_id[item_id] for item_id in ordered_ids]

            resume_text = "\n".join(item["content"] for item in evidence)
            resume_folded = resume_text.casefold()
            job_keywords = self._keywords(f"{job.title}\n{job.description}")
            supported_keywords = [
                keyword for keyword in job_keywords if keyword.casefold() in resume_folded
            ]
            skill_source = "\n".join(
                item["content"] for item in evidence
                if self._category(item["section"]) == "skills"
            ).casefold()
            supported_skills = [
                keyword for keyword in supported_keywords
                if keyword.casefold() in skill_source
            ]
            provider_skills = proposal.get("matched_skills", [])
            for skill in provider_skills if isinstance(provider_skills, list) else []:
                if (
                    isinstance(skill, str)
                    and skill.casefold() in resume_folded
                    and skill.casefold() in job.description.casefold()
                    and skill.casefold() not in {item.casefold() for item in supported_skills}
                ):
                    supported_skills.append(skill)
            missing = [
                keyword for keyword in job_keywords
                if keyword.casefold() not in resume_folded
            ]
            provider_missing = proposal.get("missing_keywords", [])
            for keyword in provider_missing if isinstance(provider_missing, list) else []:
                if (
                    isinstance(keyword, str)
                    and keyword.casefold() in job.description.casefold()
                    and keyword.casefold() not in resume_folded
                    and keyword.casefold() not in {item.casefold() for item in missing}
                ):
                    missing.append(keyword)

            buckets: dict[str, list[dict[str, Any]]] = {
                "summary": [], "experience": [], "projects": [], "education": []
            }
            for item in ordered_evidence:
                category = self._category(item["section"])
                if category == "skills":
                    continue
                buckets[category].append(
                    {
                        "text": item["content"],
                        "evidence_ids": [item["id"]],
                        "action": "调整顺序" if item["id"] in valid_order else "保留",
                    }
                )
            summary_ids = proposal.get("summary_evidence_ids", [])
            summary_evidence = [
                by_id[item_id]["content"] for item_id in summary_ids
                if isinstance(item_id, str) and item_id in by_id
            ][:2]
            if not summary_evidence:
                summary_evidence = [
                    item["content"] for item in ordered_evidence
                    if self._category(item["section"]) == "summary"
                ][:2]
            if not summary_evidence and ordered_evidence:
                summary_evidence = [ordered_evidence[0]["content"]]

            resume.summary = "\n".join(summary_evidence)
            resume.skills_json = supported_skills[:MAX_KEYWORDS]
            resume.experience_json = buckets["experience"]
            resume.projects_json = buckets["projects"]
            resume.education_json = buckets["education"]
            resume.evidence_json = evidence
            resume.warnings_json = [
                {
                    "code": "MISSING_UNSUPPORTED_KEYWORD",
                    "keyword": keyword,
                    "message": f"原始简历没有证据支持“{keyword}”，未写入定制简历。",
                    "action": "缺失",
                }
                for keyword in missing[:MAX_KEYWORDS]
            ]
            resume.generated_content_json = {
                "original": {
                    "summary": "\n".join(
                        item["content"] for item in evidence
                        if self._category(item["section"]) == "summary"
                    ),
                    "skills": self._keywords(resume_text),
                    "experience": [
                        item["content"] for item in evidence
                        if self._category(item["section"]) == "experience"
                    ],
                    "projects": [
                        item["content"] for item in evidence
                        if self._category(item["section"]) == "projects"
                    ],
                },
                "keyword_coverage": {
                    "covered": supported_keywords[:MAX_KEYWORDS],
                    "missing": missing[:MAX_KEYWORDS],
                },
                "generation_policy": "evidence-only-v1",
                "no_external_actions": True,
            }
            resume.generation_key = hashlib.sha256(
                (str(job.id) + job.description + resume_text).encode("utf-8")
            ).hexdigest()
            resume.status = TailoredResumeStatus.GENERATED
            resume.is_generating = False
            resume.error_message = None
            resume.version += 1
            self.session.commit()
            self.session.refresh(resume)
            return resume
        except Exception as exc:
            self.session.rollback()
            failed = self._get(resume_id)
            failed.status = TailoredResumeStatus.FAILED
            failed.is_generating = False
            failed.error_message = (
                getattr(exc, "public_message", None) or "定制生成失败，请重试"
            )[:500]
            failed.version += 1
            self.session.commit()
            raise

    def _assert_supported(self, resume: TailoredResume, payload: TailoredResumeUpdate) -> None:
        source = "\n".join(item["content"] for item in resume.evidence_json)
        folded = source.casefold()
        texts: list[str] = []
        if payload.summary is not None:
            texts.append(payload.summary)
        if payload.skills is not None:
            texts.extend(payload.skills)
        for group in (payload.experience, payload.projects, payload.education):
            if group is not None:
                texts.extend(item.text for item in group)
        unsupported = [text for text in texts if text and text.casefold() not in folded]
        if unsupported:
            raise UnsupportedTailoredContentError(
                "编辑内容必须逐项可追溯到原始简历，不能新增无证据事实"
            )

        evidence_by_id = {
            str(item.get("id")): item
            for item in resume.evidence_json
            if str(item.get("document_id")) == str(resume.source_document_id)
        }
        for group in (payload.experience, payload.projects, payload.education):
            for item in group or []:
                if not item.evidence_ids:
                    raise UnsupportedTailoredContentError(
                        "每条编辑内容必须引用真实的原始简历 evidence"
                    )
                for evidence_id in item.evidence_ids:
                    evidence = evidence_by_id.get(evidence_id)
                    if (
                        evidence is None
                        or str(evidence.get("source_id")) != evidence_id
                        or evidence.get("source_type") != "resume_chunk"
                    ):
                        raise UnsupportedTailoredContentError(
                            "evidence_id 不存在或不属于当前原始简历"
                        )
                    try:
                        chunk_id = uuid.UUID(evidence_id)
                    except ValueError as exc:
                        raise UnsupportedTailoredContentError(
                            "evidence_id 格式无效"
                        ) from exc
                    chunk = self.documents.get_chunk_for_document(
                        chunk_id, resume.source_document_id
                    )
                    if chunk is None:
                        raise UnsupportedTailoredContentError(
                            "evidence 对应的原始 chunk 不存在"
                        )
                    chunk_content = redact_sensitive_text(chunk.content).casefold()
                    if item.text.casefold() not in chunk_content:
                        raise UnsupportedTailoredContentError(
                            "编辑内容与所引用的 evidence chunk 不匹配"
                        )

    def update(
        self, resume_id: uuid.UUID, payload: TailoredResumeUpdate
    ) -> TailoredResume:
        resume = self._get(resume_id)
        if resume.status == TailoredResumeStatus.FINALIZED:
            self.session.rollback()
            raise ResourceConflictError("finalized tailored resume cannot be edited")
        if resume.status not in (
            TailoredResumeStatus.GENERATED, TailoredResumeStatus.EDITING
        ):
            self.session.rollback()
            raise ResourceConflictError("generate the tailored resume before editing")
        self._assert_supported(resume, payload)
        updates = payload.model_dump(exclude_unset=True)
        mapping = {
            "skills": "skills_json",
            "experience": "experience_json",
            "projects": "projects_json",
            "education": "education_json",
        }
        edited = dict(resume.user_edited_content_json)
        for field, value in updates.items():
            normalized = value
            setattr(resume, mapping.get(field, field), normalized)
            edited[field] = normalized
        resume.user_edited_content_json = edited
        resume.status = TailoredResumeStatus.EDITING
        resume.version += 1
        self.session.commit()
        self.session.refresh(resume)
        return resume

    def finalize(self, resume_id: uuid.UUID) -> TailoredResume:
        resume = self._get(resume_id)
        if resume.status == TailoredResumeStatus.FINALIZED:
            self.session.commit()
            return resume
        if resume.status not in (
            TailoredResumeStatus.GENERATED, TailoredResumeStatus.EDITING
        ):
            self.session.rollback()
            raise ResourceConflictError("generate the tailored resume before finalizing")
        resume.status = TailoredResumeStatus.FINALIZED
        resume.finalized_at = datetime.now(timezone.utc)
        resume.version += 1
        self.session.commit()
        self.session.refresh(resume)
        return resume

    def delete(self, resume_id: uuid.UUID) -> None:
        resume = self._get(resume_id)
        if resume.status == TailoredResumeStatus.FINALIZED:
            self.session.rollback()
            raise ResourceConflictError("finalized tailored resume cannot be deleted directly")
        self.session.delete(resume)
        self.session.commit()

    def export_docx(self, resume_id: uuid.UUID) -> tuple[str, bytes]:
        resume = self._get(resume_id)
        if resume.status not in (
            TailoredResumeStatus.GENERATED,
            TailoredResumeStatus.EDITING,
            TailoredResumeStatus.FINALIZED,
        ):
            self.session.rollback()
            raise ResourceConflictError("generate the tailored resume before export")
        document = DocxDocument()
        document.add_heading(resume.title, level=0)
        if resume.summary:
            document.add_heading("职业摘要", level=1)
            document.add_paragraph(resume.summary)
        if resume.skills_json:
            document.add_heading("技能", level=1)
            document.add_paragraph(" · ".join(resume.skills_json))
        for heading, items in (
            ("工作经历", resume.experience_json),
            ("项目经历", resume.projects_json),
            ("教育经历", resume.education_json),
        ):
            if items:
                document.add_heading(heading, level=1)
                for item in items:
                    document.add_paragraph(item["text"], style="List Bullet")
        document.add_paragraph(
            "本文件由证据约束的定制流程生成；未自动投递或发送。",
            style=None,
        )
        output = io.BytesIO()
        document.save(output)
        self.session.commit()
        safe_name = re.sub(r'[\\/:*?"<>|]+', "-", resume.title).strip() or "tailored-resume"
        return f"{safe_name}.docx", output.getvalue()
