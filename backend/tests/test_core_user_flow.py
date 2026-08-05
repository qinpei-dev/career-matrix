"""Minimal closed-loop tests for the core user workflow."""

from __future__ import annotations

import json
import uuid
from collections.abc import Generator
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker

from backend.app.api.dependencies import (
    get_embedding_provider,
    get_resume_tailoring_provider,
)
from backend.app.application.analysis_service import AnalysisService
from backend.app.core.config import get_settings
from backend.app.infrastructure.database.base import Base
from backend.app.infrastructure.database.models import (
    AnalysisTask,
    CandidateProfile,
    Document,
    DocumentChunk,
    TailoredResume,
)
from backend.app.infrastructure.database.session import (
    create_database_engine,
    create_session_factory,
    get_db_session,
)
from backend.app.infrastructure.llm.parser import ExtractedAnalysis, build_final_analysis
from backend.app.main import app


class CoreFlowEmbeddingProvider:
    dimension = 1024

    def embed_texts(self, texts: list[str]) -> list[list[float]]:
        return [[1.0] + [0.0] * (self.dimension - 1) for _ in texts]

    def embed_query(self, query: str) -> list[float]:
        return [1.0] + [0.0] * (self.dimension - 1)


class CoreFlowTailoringProvider:
    def tailor_resume(
        self, job_title: str, job_description: str, resume_evidence: str
    ) -> str:
        evidence = json.loads(resume_evidence)
        return json.dumps(
            {
                "summary_evidence_ids": [evidence[0]["id"]],
                "evidence_order": [item["id"] for item in evidence],
                "matched_skills": ["Python", "FastAPI"],
                "missing_keywords": [],
            }
        )


@pytest.fixture
def core_flow_api(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> Generator[tuple[TestClient, sessionmaker[Session]], None, None]:
    engine = create_database_engine(
        f"sqlite+pysqlite:///{(tmp_path / 'core-flow.db').as_posix()}"
    )
    Base.metadata.create_all(engine)
    session_factory = create_session_factory(engine)
    monkeypatch.setenv("DOCUMENT_STORAGE_PATH", str(tmp_path / "uploads"))
    get_settings.cache_clear()

    def override_session() -> Generator[Session, None, None]:
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_db_session] = override_session
    app.dependency_overrides[get_embedding_provider] = CoreFlowEmbeddingProvider
    app.dependency_overrides[get_resume_tailoring_provider] = (
        CoreFlowTailoringProvider
    )
    try:
        with TestClient(
            app, headers={"Authorization": "Bearer test-token-demo"}
        ) as client:
            yield client, session_factory
    finally:
        app.dependency_overrides.clear()
        get_settings.cache_clear()
        engine.dispose()


def _resume_bytes() -> bytes:
    document = DocxDocument()
    document.add_heading("Summary", level=1)
    document.add_paragraph("Backend engineer building reliable Python APIs.")
    document.add_heading("Skills", level=1)
    document.add_paragraph("Python, FastAPI, PostgreSQL")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _create_profile(client: TestClient) -> dict:
    response = client.post(
        "/api/v1/profiles",
        json={
            "name": "Core Flow Candidate",
            "target_role": "Backend Engineer",
            "summary": "Builds reliable backend services.",
            "skills": ["Python", "FastAPI"],
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def _create_job(client: TestClient) -> dict:
    response = client.post(
        "/api/v1/jobs",
        json={
            "title": "Backend Engineer",
            "company": "Example Co",
            "description": "Requires Python, FastAPI, and PostgreSQL.",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def _create_analysis_task(client: TestClient, job_id: str) -> dict:
    response = client.post("/api/v1/analysis-tasks", json={"job_id": job_id})
    assert response.status_code == 201, response.text
    return response.json()


def _extracted_analysis() -> ExtractedAnalysis:
    return ExtractedAnalysis.model_validate(
        {
            "job_requirements": {"core_skills": ["Python", "FastAPI"]},
            "matched_skills": ["Python", "FastAPI"],
            "missing_skills": [],
            "reasoning": ["Candidate profile contains the required skills."],
            "confidence": 0.95,
        }
    )


def _install_successful_analysis_stub(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        AnalysisService, "extract_job", lambda *_args: _extracted_analysis()
    )
    monkeypatch.setattr(
        AnalysisService,
        "analyze_job",
        lambda *_args: build_final_analysis(_extracted_analysis()),
    )


def test_t1_upload_chunks_and_profile_are_saved(core_flow_api) -> None:
    client, session_factory = core_flow_api

    uploaded = client.post(
        "/api/v1/documents/upload",
        files={"file": ("resume.docx", _resume_bytes())},
    )
    assert uploaded.status_code == 201, uploaded.text
    document_payload = uploaded.json()
    assert document_payload["id"]
    assert document_payload["status"] == "ready"
    assert len(document_payload["chunks"]) > 0

    profile_payload = _create_profile(client)
    loaded_profile = client.get("/api/v1/profiles/me")
    assert loaded_profile.status_code == 200, loaded_profile.text
    assert loaded_profile.json()["id"] == profile_payload["id"]

    with session_factory() as session:
        stored_document = session.get(Document, uuid.UUID(document_payload["id"]))
        stored_profile = session.get(CandidateProfile, uuid.UUID(profile_payload["id"]))
        assert stored_document is not None
        stored_chunks = session.scalars(
            select(DocumentChunk).where(
                DocumentChunk.document_id == stored_document.id
            )
        ).all()
        assert stored_document.status == "ready"
        assert len(stored_chunks) > 0
        assert stored_profile is not None


def test_t2_analysis_task_state_survives_a_new_database_session(
    core_flow_api, monkeypatch: pytest.MonkeyPatch
) -> None:
    client, session_factory = core_flow_api
    _create_profile(client)
    job = _create_job(client)
    task = _create_analysis_task(client, job["id"])
    _install_successful_analysis_stub(monkeypatch)

    run_response = client.post(f"/api/v1/analysis-tasks/{task['task_id']}/run")
    assert run_response.status_code == 200, run_response.text
    run_payload = run_response.json()
    assert run_payload["status"] == "WAITING_FOR_REVIEW"

    task_id = uuid.UUID(task["task_id"])
    with session_factory() as committed_session:
        committed = committed_session.get(AnalysisTask, task_id)
        assert committed is not None
        committed_session.commit()

    with session_factory() as restored_session:
        restored = restored_session.get(AnalysisTask, task_id)
        assert restored is not None
        assert restored.status == run_payload["status"]
        assert restored.result_id is not None
        assert restored.result_payload
        assert restored.error_code is None
        assert restored.error_message is None


def test_t3_llm_timeout_fails_cleanly_and_can_retry(
    core_flow_api, monkeypatch: pytest.MonkeyPatch
) -> None:
    client, session_factory = core_flow_api
    _create_profile(client)
    job = _create_job(client)
    task = _create_analysis_task(client, job["id"])
    calls = 0

    def timeout_once(*_args):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise TimeoutError("simulated LLM timeout")
        return _extracted_analysis()

    monkeypatch.setattr(AnalysisService, "extract_job", timeout_once)
    monkeypatch.setattr(
        AnalysisService,
        "analyze_job",
        lambda *_args: build_final_analysis(_extracted_analysis()),
    )

    failed_response = client.post(
        f"/api/v1/analysis-tasks/{task['task_id']}/run"
    )
    assert failed_response.status_code == 200, failed_response.text
    failed = failed_response.json()
    assert failed["status"] == "FAILED"
    assert failed["current_step"] == "ANALYZING"
    assert failed["error_code"] == "TimeoutError"
    assert failed["error_message"]
    assert failed["retry_count"] == 0

    with session_factory() as session:
        stored_failure = session.get(AnalysisTask, uuid.UUID(task["task_id"]))
        assert stored_failure is not None
        assert stored_failure.status == "FAILED"
        assert stored_failure.error_message
        assert stored_failure.is_running is False

    retry_response = client.post(
        f"/api/v1/analysis-tasks/{task['task_id']}/retry"
    )
    assert retry_response.status_code == 200, retry_response.text
    retried = retry_response.json()
    assert retried["status"] == "WAITING_FOR_REVIEW"
    assert retried["retry_count"] == 1
    assert retried["retry_count"] <= retried["max_retries"]
    assert calls == 2


def test_t4_generated_result_persists_and_exports_to_a_nonempty_file(
    core_flow_api, tmp_path: Path
) -> None:
    client, session_factory = core_flow_api
    uploaded = client.post(
        "/api/v1/documents/upload",
        files={"file": ("resume.docx", _resume_bytes())},
    )
    assert uploaded.status_code == 201, uploaded.text
    job = _create_job(client)

    created = client.post(
        "/api/v1/tailored-resumes",
        json={
            "job_id": job["id"],
            "source_document_id": uploaded.json()["id"],
        },
    )
    assert created.status_code == 201, created.text
    resume_id = created.json()["tailored_resume_id"]

    generated_response = client.post(
        f"/api/v1/tailored-resumes/{resume_id}/generate"
    )
    assert generated_response.status_code == 200, generated_response.text
    generated = generated_response.json()
    assert generated["status"] == "GENERATED"
    assert generated["generated_content_json"]

    queried = client.get(f"/api/v1/tailored-resumes/{resume_id}")
    assert queried.status_code == 200, queried.text
    assert queried.json()["generated_content_json"] == generated[
        "generated_content_json"
    ]

    with session_factory() as restored_session:
        stored = restored_session.get(TailoredResume, uuid.UUID(resume_id))
        assert stored is not None
        assert stored.generated_content_json

    exported = client.get(f"/api/v1/tailored-resumes/{resume_id}/export.docx")
    assert exported.status_code == 200, exported.text
    export_path = tmp_path / "tailored-resume.docx"
    export_path.write_bytes(exported.content)
    assert export_path.is_file()
    assert export_path.stat().st_size > 0
